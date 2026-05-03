from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader


@dataclass
class ParsedBlock:
    section: str
    page: int
    text: str
    table_markdown: str | None = None
    parser_status: str = "ok"
    parser_confidence: float = 0.95


def parse_document(path: Path) -> list[ParsedBlock]:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return parse_docx(path)
    if suffix == ".pdf":
        return parse_pdf(path)
    raise ValueError(f"Unsupported document type: {path.suffix}")


def parse_docx(path: Path) -> list[ParsedBlock]:
    doc = DocxDocument(path)
    blocks: list[ParsedBlock] = []
    current_section = "Front Matter"
    current_text: list[str] = []

    def flush() -> None:
        nonlocal current_text
        text = " ".join(item.strip() for item in current_text if item.strip()).strip()
        if text:
            blocks.append(ParsedBlock(section=current_section, page=1, text=text))
        current_text = []

    for paragraph in doc.paragraphs:
        value = paragraph.text.strip()
        if not value:
            continue
        style = paragraph.style.name.lower() if paragraph.style and paragraph.style.name else ""
        if "heading 1" in style:
            flush()
            current_section = value
        elif "title" in style or "heading 0" in style:
            continue
        else:
            current_text.append(value)
    flush()

    for table in doc.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        if not rows:
            continue
        markdown = _table_to_markdown(rows)
        section = "Readiness Review Table" if "readiness_pct" in rows[0] else current_section
        blocks.append(
            ParsedBlock(
                section=section,
                page=1,
                text=f"Table extracted from {section}.",
                table_markdown=markdown,
                parser_confidence=0.9,
            )
        )
    return blocks


def parse_pdf(path: Path) -> list[ParsedBlock]:
    reader = PdfReader(str(path))
    blocks: list[ParsedBlock] = []
    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        sections = text.split("SECTION:")
        front = sections[0].strip()
        if front and not blocks:
            blocks.append(ParsedBlock(section="Front Matter", page=page_number, text=front, parser_confidence=0.88))
        for raw_section in sections[1:]:
            lines = [line.strip() for line in raw_section.splitlines() if line.strip()]
            if not lines:
                continue
            section = lines[0]
            body = " ".join(lines[1:]).strip()
            blocks.append(ParsedBlock(section=section, page=page_number, text=body, parser_confidence=0.88))
    return blocks


def _table_to_markdown(rows: list[list[str]]) -> str:
    header = rows[0]
    separator = ["---"] * len(header)
    body = rows[1:]
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(separator) + " |",
    ]
    for row in body:
        lines.append("| " + " | ".join(row) + " |")
    return "\n".join(lines)
