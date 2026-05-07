from __future__ import annotations

import csv
import hashlib
from dataclasses import dataclass
from pathlib import Path

import yaml
from docx import Document as DocxDocument
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from defence_agent.config import get_settings


@dataclass(frozen=True)
class SyntheticDocSpec:
    doc_id: str
    filename: str
    title: str
    doc_family: str
    version: str
    status: str
    language: str
    access_level: str
    owner: str
    effective_date: str
    review_due: str
    source_type: str
    sections: tuple[tuple[str, int, str], ...]
    authoritative_rank: int = 100
    supersedes: tuple[str, ...] = ()
    superseded_by: tuple[str, ...] = ()
    cross_references: tuple[str, ...] = ()
    applies_to: tuple[str, ...] = ()
    not_applicable_to: tuple[str, ...] = ()

    @property
    def markdown_filename(self) -> str:
        return f"{self.doc_id}.md"


SPECS: tuple[SyntheticDocSpec, ...] = (
    SyntheticDocSpec(
        doc_id="PB-SOP-2025",
        filename="PB-SOP-2025.docx",
        title="Planning Brief Approval SOP",
        doc_family="planning_brief_approval",
        version="2025.1",
        status="approved",
        language="en",
        access_level="public_internal",
        owner="Planning Policy Directorate",
        effective_date="2025-02-01",
        review_due="2026-03-15",
        source_type="docx",
        authoritative_rank=10,
        supersedes=("PB-SOP-2024",),
        cross_references=("PB-CHK-2025", "LOG-RET-2025", "ANNEX-HANDLING-2025"),
        applies_to=("planning_brief", "restricted_annex", "public_release"),
        sections=(
            (
                "Current Approval Process",
                1,
                "The current approved planning brief approval process requires intake and purpose classification, evidence pack review, planning brief drafting, section chief review within 2 business days, legal or policy review when the brief includes external distribution, restricted annexes, or cross-departmental commitments, director approval before circulation, and archival of the evidence and decision log.",
            ),
            (
                "Urgent Exception Rule",
                2,
                "Urgency may compress the order of sign-offs, but it cannot skip evidence review. The rationale must be recorded, and retrospective director confirmation is required within 1 business day.",
            ),
            (
                "Evidence Pack Cross-Reference",
                2,
                "Before review, the required evidence pack must satisfy PB-CHK-2025.",
            ),
        ),
    ),
    SyntheticDocSpec(
        doc_id="PB-SOP-2024",
        filename="PB-SOP-2024.docx",
        title="Planning Brief Approval SOP",
        doc_family="planning_brief_approval",
        version="2024.3",
        status="superseded",
        language="en",
        access_level="public_internal",
        owner="Planning Policy Directorate",
        effective_date="2024-01-10",
        review_due="2025-01-10",
        source_type="docx",
        authoritative_rank=60,
        superseded_by=("PB-SOP-2025",),
        applies_to=("planning_brief",),
        sections=(
            (
                "Superseded Review Process",
                1,
                "The superseded 2024 planning brief process allowed section chief review within 3 business days. Legal and policy review was required only for public release.",
            ),
            (
                "Urgent Provisional Circulation",
                2,
                "Urgent provisional circulation was allowed after duty officer approval. Evidence review was recommended but not mandatory before urgent provisional circulation.",
            ),
        ),
    ),
    SyntheticDocSpec(
        doc_id="PB-SOP-2026-DRAFT",
        filename="PB-SOP-2026-DRAFT.docx",
        title="Planning Brief Approval SOP",
        doc_family="planning_brief_approval",
        version="2026.0-draft",
        status="draft",
        language="en",
        access_level="public_internal",
        owner="Planning Policy Directorate",
        effective_date="2026-04-01",
        review_due="2026-12-31",
        source_type="docx",
        authoritative_rank=90,
        supersedes=(),
        superseded_by=(),
        applies_to=("planning_brief", "draft_policy"),
        not_applicable_to=("current_approved_guidance",),
        sections=(
            (
                "Draft Pilot Language",
                1,
                "Draft only. This pilot proposes self-approval for low-risk briefs after automated checklist completion. This language is not approved guidance and must not be used for current approved procedure questions.",
            ),
        ),
    ),
    SyntheticDocSpec(
        doc_id="PB-CHK-2025",
        filename="PB-CHK-2025.pdf",
        title="Planning Brief Evidence Checklist",
        doc_family="planning_brief_checklist",
        version="2025.2",
        status="approved",
        language="en",
        access_level="public_internal",
        owner="Planning Policy Directorate",
        effective_date="2025-02-01",
        review_due="2026-08-01",
        source_type="pdf",
        authoritative_rank=15,
        cross_references=("PB-SOP-2025", "LOG-RET-2025"),
        applies_to=("planning_brief", "evidence_pack"),
        sections=(
            (
                "Required Evidence Pack",
                1,
                "Before review, the evidence pack must include a problem statement, policy basis, source list, assumptions, impacted stakeholders, risk rating, options considered, recommendation, citation table, open questions, and decision log stub.",
            ),
        ),
    ),
    SyntheticDocSpec(
        doc_id="EC-PROC-2025",
        filename="EC-PROC-2025.pdf",
        title="Emergency Communications Procedure",
        doc_family="emergency_communications",
        version="2025.1",
        status="approved",
        language="en",
        access_level="public_internal",
        owner="Emergency Coordination Office",
        effective_date="2025-06-01",
        review_due="2026-04-01",
        source_type="pdf",
        authoritative_rank=12,
        cross_references=("PR-GUIDE-2025", "LOG-RET-2025"),
        applies_to=("emergency_communications", "public_release"),
        sections=(
            (
                "Activation And Approval Gates",
                1,
                "The emergency communications procedure covers activation conditions, approval gates, and required evidence. Public dissemination requires citation to the reporting guide.",
            ),
            (
                "Timelines",
                1,
                "The initial acknowledgement is due within 15 minutes, the operational situation update within 45 minutes, the executive summary within 90 minutes, and the final record within 1 business day.",
            ),
            (
                "Required Evidence",
                2,
                "Required evidence includes source event log, responsible desk, time of receipt, confidence level, distribution list, and approving official.",
            ),
        ),
    ),
    SyntheticDocSpec(
        doc_id="EC-PROC-2025-FR",
        filename="EC-PROC-2025-FR.pdf",
        title="Procedure de communications d'urgence",
        doc_family="emergency_communications",
        version="2025.1-fr",
        status="approved",
        language="fr",
        access_level="public_internal",
        owner="Bureau de coordination des urgences",
        effective_date="2025-06-01",
        review_due="2026-04-01",
        source_type="pdf",
        authoritative_rank=12,
        cross_references=("PR-GUIDE-2025-FR",),
        applies_to=("emergency_communications", "public_release"),
        sections=(
            (
                "Delais",
                1,
                "La procedure de communications d'urgence exige un accuse de reception initial dans les 15 minutes, une mise a jour operationnelle dans les 45 minutes, un resume executif dans les 90 minutes, et un dossier final dans un jour ouvrable.",
            ),
        ),
    ),
    SyntheticDocSpec(
        doc_id="ANNEX-HANDLING-2025",
        filename="ANNEX-HANDLING-2025.pdf",
        title="Restricted Annex Handling Guide",
        doc_family="annex_handling",
        version="2025.1",
        status="approved",
        language="en",
        access_level="restricted",
        owner="Records and Security Office",
        effective_date="2025-07-01",
        review_due="2026-06-30",
        source_type="pdf",
        authoritative_rank=8,
        cross_references=("CLASS-MARK-2025", "CLASS-ANNEX-B-2025"),
        applies_to=("restricted_annex", "external_distribution"),
        sections=(
            (
                "External Distribution Control",
                1,
                "Restricted annex handling before external distribution requires access confirmation, lead approval, distribution minimization, annex marking validation, and an evidence log entry. This is synthetic safe content.",
            ),
            (
                "Disclosure Boundary",
                1,
                "Restricted annex handling details must not be shown to users without restricted access and an authorized role.",
            ),
        ),
    ),
    SyntheticDocSpec(
        doc_id="LOG-RET-2025",
        filename="LOG-RET-2025.docx",
        title="Evidence and Decision Log Retention Procedure",
        doc_family="evidence_log_retention",
        version="2025.1",
        status="approved",
        language="en",
        access_level="public_internal",
        owner="Records Management Office",
        effective_date="2025-01-15",
        review_due="2026-01-30",
        source_type="docx",
        authoritative_rank=20,
        applies_to=("audit_log", "traceability", "decision_record"),
        sections=(
            (
                "Evidence And Decision Logs",
                1,
                "Evidence logs must record source title, source version, section or page, reviewer, and decision timestamp. Decision logs must be retained for 7 years.",
            ),
            (
                "Traceability Requirement",
                1,
                "The log must distinguish facts from assumptions. Each briefing decision must include a trace ID.",
            ),
        ),
    ),
    SyntheticDocSpec(
        doc_id="SCANNED-MANUAL-EXCERPT-1999",
        filename="SCANNED-MANUAL-EXCERPT-1999.pdf",
        title="Legacy Field Manual Excerpt",
        doc_family="legacy_manual",
        version="1999-scan",
        status="superseded",
        language="en",
        access_level="public_internal",
        owner="Archives",
        effective_date="1999-03-01",
        review_due="2000-03-01",
        source_type="scanned_pdf",
        authoritative_rank=95,
        superseded_by=("PB-SOP-2025", "SUPERSESSION-BULLETIN-2025"),
        applies_to=("legacy_manual",),
        not_applicable_to=("current_approved_guidance",),
        sections=(
            (
                "OCR Simulated Legacy Process",
                1,
                "OCR simulated text: LEG4CY MANUAL EXCERPT. Approval was paper-based and routed by courier. This old process is superseded and should not answer current approved procedure questions.",
            ),
        ),
    ),
    SyntheticDocSpec(
        doc_id="JPD-2025",
        filename="JPD-2025.docx",
        title="Joint Planning Doctrine",
        doc_family="joint_planning_doctrine",
        version="2025.1",
        status="approved",
        language="en",
        access_level="public_internal",
        owner="Joint Planning Office",
        effective_date="2025-03-01",
        review_due="2026-09-30",
        source_type="docx",
        authoritative_rank=9,
        supersedes=("JPD-2024", "FIELD-MANUAL-2008-SCAN"),
        cross_references=("PB-SOP-2025", "PB-CHK-2025", "LOG-RET-2025", "READINESS-DIR-2025"),
        applies_to=("interagency_planning", "emergency_brief", "director_approval"),
        sections=(
            (
                "Joint Planning Control Points",
                1,
                "JPD-2025 requires planners to classify the planning purpose, identify interagency dependencies, confirm readiness evidence, and prepare a director-ready planning brief before action is recommended.",
            ),
            (
                "Required Cross-References",
                2,
                "Before director approval, JPD-2025 requires PB-CHK-2025 evidence, LOG-RET-2025 trace logging, and READINESS-DIR-2025 readiness evidence. Public-release planning must also follow PR-GUIDE-2025.",
            ),
            (
                "Currentness Rule",
                3,
                "JPD-2025 supersedes JPD-2024 and the FIELD-MANUAL-2008-SCAN excerpt for current approved joint planning guidance.",
            ),
        ),
    ),
    SyntheticDocSpec(
        doc_id="JPD-2024",
        filename="JPD-2024.docx",
        title="Joint Planning Doctrine",
        doc_family="joint_planning_doctrine",
        version="2024.2",
        status="superseded",
        language="en",
        access_level="public_internal",
        owner="Joint Planning Office",
        effective_date="2024-02-01",
        review_due="2025-02-01",
        source_type="docx",
        authoritative_rank=65,
        superseded_by=("JPD-2025",),
        applies_to=("interagency_planning",),
        not_applicable_to=("current_approved_guidance",),
        sections=(
            (
                "Superseded Coordination Model",
                1,
                "The 2024 doctrine allowed readiness confirmation after director review when time was limited. This sequence is superseded by JPD-2025.",
            ),
            (
                "Superseded Evidence Rule",
                2,
                "The 2024 doctrine recommended but did not require a complete evidence pack before director review for interagency planning briefs.",
            ),
        ),
    ),
    SyntheticDocSpec(
        doc_id="JPD-2026-DRAFT",
        filename="JPD-2026-DRAFT.docx",
        title="Joint Planning Doctrine",
        doc_family="joint_planning_doctrine",
        version="2026.0-draft",
        status="draft",
        language="en",
        access_level="public_internal",
        owner="Joint Planning Office",
        effective_date="2026-04-15",
        review_due="2026-12-31",
        source_type="docx",
        authoritative_rank=92,
        applies_to=("draft_policy",),
        not_applicable_to=("current_approved_guidance",),
        sections=(
            (
                "Draft Pilot Concepts",
                1,
                "Draft only. The 2026 draft proposes automated readiness attestation for low-risk interagency briefs. It is not approved and must not answer current approved guidance questions.",
            ),
        ),
    ),
    SyntheticDocSpec(
        doc_id="IC-MOU-2025",
        filename="IC-MOU-2025.pdf",
        title="Interagency Coordination Memorandum of Understanding",
        doc_family="interagency_coordination",
        version="2025.1",
        status="approved",
        language="en",
        access_level="public_internal",
        owner="Interagency Coordination Secretariat",
        effective_date="2025-04-01",
        review_due="2026-10-31",
        source_type="pdf",
        authoritative_rank=14,
        cross_references=("IC-CONOPS-2025", "PR-GUIDE-2025", "LOG-RET-2025"),
        applies_to=("interagency_planning", "public_release"),
        sections=(
            (
                "Coordination Commitments",
                1,
                "The MOU requires a lead desk, named partner contacts, distribution constraints, and agreement on what information can be shared before an interagency planning brief is circulated.",
            ),
            (
                "Public Release Boundary",
                2,
                "If interagency material may be released publicly, the lead desk must apply PR-GUIDE-2025 and record the decision under LOG-RET-2025.",
            ),
        ),
    ),
    SyntheticDocSpec(
        doc_id="IC-CONOPS-2025",
        filename="IC-CONOPS-2025.docx",
        title="Interagency Emergency Planning CONOPS",
        doc_family="interagency_coordination",
        version="2025.1",
        status="approved",
        language="en",
        access_level="public_internal",
        owner="Interagency Coordination Secretariat",
        effective_date="2025-04-15",
        review_due="2026-11-30",
        source_type="docx",
        authoritative_rank=16,
        cross_references=("EC-PROC-2025", "JPD-2025", "IC-MOU-2025"),
        applies_to=("interagency_planning", "emergency_communications"),
        sections=(
            (
                "Emergency Planning Workflow",
                1,
                "For interagency emergencies, staff use EC-PROC-2025 for communications timelines, JPD-2025 for planning control points, and IC-MOU-2025 for partner coordination rules.",
            ),
            (
                "Decision Record",
                2,
                "The CONOPS requires the decision record to name the lead desk, partner agencies, response clock start time, release boundary, and trace ID.",
            ),
        ),
    ),
    SyntheticDocSpec(
        doc_id="IC-ANNEX-DATA-2025",
        filename="IC-ANNEX-DATA-2025.pdf",
        title="Interagency Data Sharing Annex",
        doc_family="interagency_data_sharing",
        version="2025.1",
        status="approved",
        language="en",
        access_level="restricted",
        owner="Interagency Coordination Secretariat",
        effective_date="2025-05-01",
        review_due="2026-11-30",
        source_type="pdf",
        authoritative_rank=7,
        cross_references=("ANNEX-HANDLING-2025", "CLASS-ANNEX-B-2025"),
        applies_to=("restricted_annex", "interagency_data"),
        sections=(
            (
                "Restricted Data Handling",
                1,
                "Synthetic restricted content. Interagency data-sharing annex material requires need-to-know validation, minimized excerpts, restricted marking checks, and an audit entry before external partner distribution.",
            ),
            (
                "Analyst Boundary",
                2,
                "Planning analysts without restricted access must receive only a statement that restricted data-sharing detail is unavailable to their persona.",
            ),
        ),
    ),
    SyntheticDocSpec(
        doc_id="PR-GUIDE-2025",
        filename="PR-GUIDE-2025.pdf",
        title="Public Release and Disclosure Guide",
        doc_family="public_release",
        version="2025.1",
        status="approved",
        language="en",
        access_level="public_internal",
        owner="Public Affairs and Disclosure Office",
        effective_date="2025-03-20",
        review_due="2026-09-15",
        source_type="pdf",
        authoritative_rank=11,
        supersedes=("PR-GUIDE-2023",),
        cross_references=("PR-CHECKLIST-2025", "CLASS-MARK-2025", "LOG-RET-2025", "EC-PROC-2025"),
        applies_to=("public_release", "emergency_communications"),
        sections=(
            (
                "Disclosure Review",
                1,
                "Public release requires source confirmation, classification marking review, personal information screening, partner coordination, and approval by the disclosure lead before publication.",
            ),
            (
                "Emergency Communications Tie-In",
                2,
                "When emergency communications may become public, staff must cite EC-PROC-2025 timelines, apply PR-CHECKLIST-2025, and retain the release decision under LOG-RET-2025.",
            ),
        ),
    ),
    SyntheticDocSpec(
        doc_id="PR-CHECKLIST-2025",
        filename="PR-CHECKLIST-2025.pdf",
        title="Public Release Checklist",
        doc_family="public_release_checklist",
        version="2025.1",
        status="approved",
        language="en",
        access_level="public_internal",
        owner="Public Affairs and Disclosure Office",
        effective_date="2025-03-20",
        review_due="2026-09-15",
        source_type="pdf",
        authoritative_rank=18,
        cross_references=("PR-GUIDE-2025", "CLASS-MARK-2025"),
        applies_to=("public_release", "checklist"),
        sections=(
            (
                "Release Checklist Items",
                1,
                "The release checklist requires document title, source version, classification marking, disclosure lead approval, partner notification, public summary, and retained release decision.",
            ),
        ),
    ),
    SyntheticDocSpec(
        doc_id="PR-GUIDE-2023",
        filename="PR-GUIDE-2023.pdf",
        title="Public Release and Disclosure Guide",
        doc_family="public_release",
        version="2023.4",
        status="superseded",
        language="en",
        access_level="public_internal",
        owner="Public Affairs and Disclosure Office",
        effective_date="2023-05-01",
        review_due="2024-05-01",
        source_type="pdf",
        authoritative_rank=70,
        superseded_by=("PR-GUIDE-2025",),
        applies_to=("public_release",),
        not_applicable_to=("current_approved_guidance",),
        sections=(
            (
                "Superseded Release Review",
                1,
                "The 2023 guide allowed public release after a public affairs review only. PR-GUIDE-2025 supersedes this rule and adds classification, partner, and retention checks.",
            ),
        ),
    ),
    SyntheticDocSpec(
        doc_id="PR-GUIDE-2025-FR",
        filename="PR-GUIDE-2025-FR.pdf",
        title="Guide de diffusion publique et de divulgation",
        doc_family="public_release",
        version="2025.1-fr",
        status="approved",
        language="fr",
        access_level="public_internal",
        owner="Bureau des affaires publiques et de la divulgation",
        effective_date="2025-03-20",
        review_due="2026-09-15",
        source_type="pdf",
        authoritative_rank=11,
        cross_references=("PR-CHECKLIST-2025", "CLASS-MARK-2025"),
        applies_to=("public_release", "bilingual"),
        sections=(
            (
                "Examen De Divulgation",
                1,
                "La diffusion publique exige la confirmation des sources, la verification du marquage de classification, le controle des renseignements personnels, la coordination avec les partenaires et l'approbation du responsable de la divulgation.",
            ),
        ),
    ),
    SyntheticDocSpec(
        doc_id="CLASS-MARK-2025",
        filename="CLASS-MARK-2025.docx",
        title="Classification and Marking Standard",
        doc_family="classification_marking",
        version="2025.1",
        status="approved",
        language="en",
        access_level="public_internal",
        owner="Records and Security Office",
        effective_date="2025-02-15",
        review_due="2026-07-31",
        source_type="docx",
        authoritative_rank=13,
        cross_references=("CLASS-ERRATA-2025", "CLASS-ANNEX-B-2025", "ANNEX-HANDLING-2025"),
        applies_to=("classification", "public_release", "restricted_annex"),
        sections=(
            (
                "Marking Review",
                1,
                "Before public release or external distribution, staff must confirm the document banner, portion markings, annex labels, and dissemination caveats.",
            ),
            (
                "Restricted Annex Marker",
                2,
                "Restricted annexes must be marked as restricted attachments and handled according to ANNEX-HANDLING-2025. CLASS-ERRATA-2025 corrects the review office name.",
            ),
        ),
    ),
    SyntheticDocSpec(
        doc_id="CLASS-ERRATA-2025",
        filename="CLASS-ERRATA-2025.pdf",
        title="Classification Standard Errata",
        doc_family="classification_marking",
        version="2025.1-errata",
        status="approved",
        language="en",
        access_level="public_internal",
        owner="Records and Security Office",
        effective_date="2025-05-10",
        review_due="2026-07-31",
        source_type="pdf",
        authoritative_rank=6,
        cross_references=("CLASS-MARK-2025",),
        applies_to=("classification", "conflict_resolution"),
        sections=(
            (
                "Errata Authority",
                1,
                "This errata corrects CLASS-MARK-2025: the Records and Security Office, not the Public Affairs desk, is the final authority for restricted annex marking validation.",
            ),
        ),
    ),
    SyntheticDocSpec(
        doc_id="CLASS-ANNEX-B-2025",
        filename="CLASS-ANNEX-B-2025.pdf",
        title="Classification Annex B Handling Notes",
        doc_family="classification_marking",
        version="2025.1",
        status="approved",
        language="en",
        access_level="restricted",
        owner="Records and Security Office",
        effective_date="2025-02-15",
        review_due="2026-07-31",
        source_type="pdf",
        authoritative_rank=7,
        cross_references=("CLASS-MARK-2025", "ANNEX-HANDLING-2025"),
        applies_to=("restricted_annex", "classification"),
        sections=(
            (
                "Restricted Marking Detail",
                1,
                "Synthetic restricted content. Annex B requires restricted marking validation, access-list confirmation, excerpt minimization, and trace logging before external partner use.",
            ),
        ),
    ),
    SyntheticDocSpec(
        doc_id="READINESS-DIR-2025",
        filename="READINESS-DIR-2025.docx",
        title="Readiness Evidence Directive",
        doc_family="readiness_evidence",
        version="2025.1",
        status="approved",
        language="en",
        access_level="public_internal",
        owner="Readiness Secretariat",
        effective_date="2025-03-05",
        review_due="2026-10-15",
        source_type="docx",
        authoritative_rank=17,
        cross_references=("readiness_review_table", "corrective_action_tracker", "JPD-2025"),
        applies_to=("readiness", "planning_brief", "structured_analysis"),
        sections=(
            (
                "Readiness Evidence Rule",
                1,
                "Planning briefs that depend on unit readiness must cite the readiness review table, identify units below threshold, and record mitigation owners before director approval.",
            ),
            (
                "Threshold Rule",
                2,
                "The readiness threshold is 80 percent. Units below threshold require corrective action tracking before the planning brief can be recommended.",
            ),
        ),
    ),
    SyntheticDocSpec(
        doc_id="FIELD-MANUAL-2008-SCAN",
        filename="FIELD-MANUAL-2008-SCAN.pdf",
        title="Field Planning Manual Scanned Excerpt",
        doc_family="legacy_manual",
        version="2008-scan",
        status="superseded",
        language="en",
        access_level="public_internal",
        owner="Archives",
        effective_date="2008-06-01",
        review_due="2009-06-01",
        source_type="scanned_pdf",
        authoritative_rank=96,
        superseded_by=("JPD-2025", "SUPERSESSION-BULLETIN-2025"),
        applies_to=("legacy_manual",),
        not_applicable_to=("current_approved_guidance",),
        sections=(
            (
                "OCR Noisy Approval Notes",
                1,
                "OCR simulated text: F1ELD PL4NNING MANVAL 2008. Appr0val may proceed after teleph0ne concurrence. This scanned excerpt is superseded and should be used only as legacy context.",
            ),
            (
                "OCR Artifact Caveat",
                1,
                "Text contains OCR artifacts such as PL4NNING, teleph0ne, and Appr0val. Current approved guidance is JPD-2025 and PB-SOP-2025.",
            ),
        ),
    ),
    SyntheticDocSpec(
        doc_id="SUPERSESSION-BULLETIN-2025",
        filename="SUPERSESSION-BULLETIN-2025.pdf",
        title="Doctrine Supersession Bulletin",
        doc_family="supersession_bulletin",
        version="2025.1",
        status="approved",
        language="en",
        access_level="public_internal",
        owner="Doctrine Stewardship Office",
        effective_date="2025-08-01",
        review_due="2026-12-31",
        source_type="pdf",
        authoritative_rank=5,
        cross_references=("JPD-2025", "PB-SOP-2025", "PR-GUIDE-2025", "FIELD-MANUAL-2008-SCAN"),
        applies_to=("currentness", "stale_guidance_warning"),
        sections=(
            (
                "Supersession Notice",
                1,
                "JPD-2025, PB-SOP-2025, and PR-GUIDE-2025 are current approved sources. FIELD-MANUAL-2008-SCAN, SCANNED-MANUAL-EXCERPT-1999, JPD-2024, and PR-GUIDE-2023 are superseded for current guidance.",
            ),
        ),
    ),
)


TRACKER_ROWS: tuple[tuple[str, str, str, str, str, str, str, str, str], ...] = tuple(
    (
        spec.doc_id,
        spec.title,
        spec.doc_family,
        spec.owner,
        spec.status,
        spec.effective_date,
        spec.review_due,
        spec.access_level,
        spec.language,
    )
    for spec in SPECS
)


TABLE_DEFINITIONS: dict[str, tuple[tuple[str, ...], tuple[dict[str, str], ...]]] = {
    "readiness_review_table.csv": (
        (
            "unit_id",
            "unit_name",
            "owner",
            "readiness_percent",
            "threshold_percent",
            "review_date",
            "status",
            "linked_doc_id",
            "access_level",
            "language",
        ),
        (
            {
                "unit_id": "U-100",
                "unit_name": "Central Planning Cell",
                "owner": "Readiness Secretariat",
                "readiness_percent": "92",
                "threshold_percent": "80",
                "review_date": "2026-04-20",
                "status": "ready",
                "linked_doc_id": "READINESS-DIR-2025",
                "access_level": "public_internal",
                "language": "en",
            },
            {
                "unit_id": "U-210",
                "unit_name": "Emergency Communications Desk",
                "owner": "Emergency Coordination Office",
                "readiness_percent": "76",
                "threshold_percent": "80",
                "review_date": "2026-04-22",
                "status": "below_threshold",
                "linked_doc_id": "EC-PROC-2025",
                "access_level": "public_internal",
                "language": "en",
            },
            {
                "unit_id": "U-305",
                "unit_name": "Interagency Liaison Desk",
                "owner": "Interagency Coordination Secretariat",
                "readiness_percent": "68",
                "threshold_percent": "80",
                "review_date": "2026-04-25",
                "status": "below_threshold",
                "linked_doc_id": "IC-CONOPS-2025",
                "access_level": "public_internal",
                "language": "en",
            },
            {
                "unit_id": "U-440",
                "unit_name": "Disclosure Review Team",
                "owner": "Public Affairs and Disclosure Office",
                "readiness_percent": "84",
                "threshold_percent": "80",
                "review_date": "2026-04-26",
                "status": "ready",
                "linked_doc_id": "PR-GUIDE-2025",
                "access_level": "public_internal",
                "language": "en",
            },
        ),
    ),
    "approval_register.csv": (
        (
            "approval_id",
            "brief_id",
            "title",
            "owner",
            "required_source",
            "approval_stage",
            "due_date",
            "status",
            "access_level",
            "language",
        ),
        (
            {
                "approval_id": "APR-001",
                "brief_id": "BRIEF-EMERG-26",
                "title": "Emergency Public Release Planning Brief",
                "owner": "Joint Planning Office",
                "required_source": "JPD-2025",
                "approval_stage": "director_review",
                "due_date": "2026-05-08",
                "status": "pending",
                "access_level": "public_internal",
                "language": "en",
            },
            {
                "approval_id": "APR-002",
                "brief_id": "BRIEF-EMERG-26",
                "title": "Emergency Public Release Planning Brief",
                "owner": "Public Affairs and Disclosure Office",
                "required_source": "PR-GUIDE-2025",
                "approval_stage": "disclosure_lead_review",
                "due_date": "2026-05-07",
                "status": "pending",
                "access_level": "public_internal",
                "language": "en",
            },
            {
                "approval_id": "APR-003",
                "brief_id": "BRIEF-ANNEX-26",
                "title": "Restricted Annex Planning Brief",
                "owner": "Records and Security Office",
                "required_source": "ANNEX-HANDLING-2025",
                "approval_stage": "restricted_annex_review",
                "due_date": "2026-05-07",
                "status": "restricted_pending",
                "access_level": "restricted",
                "language": "en",
            },
        ),
    ),
    "corrective_action_tracker.csv": (
        (
            "action_id",
            "unit_id",
            "unit_name",
            "owner",
            "linked_doc_id",
            "issue",
            "severity",
            "due_date",
            "status",
            "access_level",
            "language",
        ),
        (
            {
                "action_id": "CA-001",
                "unit_id": "U-210",
                "unit_name": "Emergency Communications Desk",
                "owner": "Emergency Coordination Office",
                "linked_doc_id": "EC-PROC-2025",
                "issue": "Update distribution list evidence before director approval",
                "severity": "medium",
                "due_date": "2026-05-10",
                "status": "open",
                "access_level": "public_internal",
                "language": "en",
            },
            {
                "action_id": "CA-002",
                "unit_id": "U-305",
                "unit_name": "Interagency Liaison Desk",
                "owner": "Interagency Coordination Secretariat",
                "linked_doc_id": "IC-CONOPS-2025",
                "issue": "Confirm partner contact list and release boundary",
                "severity": "high",
                "due_date": "2026-05-09",
                "status": "open",
                "access_level": "public_internal",
                "language": "en",
            },
            {
                "action_id": "CA-003",
                "unit_id": "U-305",
                "unit_name": "Interagency Liaison Desk",
                "owner": "Records and Security Office",
                "linked_doc_id": "IC-ANNEX-DATA-2025",
                "issue": "Restricted data sharing review required",
                "severity": "high",
                "due_date": "2026-05-09",
                "status": "restricted_open",
                "access_level": "restricted",
                "language": "en",
            },
        ),
    ),
    "annex_inventory.csv": (
        (
            "annex_id",
            "annex_title",
            "owner",
            "linked_doc_id",
            "brief_id",
            "classification",
            "external_distribution_allowed",
            "last_reviewed",
            "access_level",
            "language",
        ),
        (
            {
                "annex_id": "ANN-001",
                "annex_title": "Public Emergency Timeline Appendix",
                "owner": "Emergency Coordination Office",
                "linked_doc_id": "EC-PROC-2025",
                "brief_id": "BRIEF-EMERG-26",
                "classification": "public_internal",
                "external_distribution_allowed": "yes_after_disclosure_review",
                "last_reviewed": "2026-04-28",
                "access_level": "public_internal",
                "language": "en",
            },
            {
                "annex_id": "ANN-002",
                "annex_title": "Restricted Partner Data Appendix",
                "owner": "Interagency Coordination Secretariat",
                "linked_doc_id": "IC-ANNEX-DATA-2025",
                "brief_id": "BRIEF-ANNEX-26",
                "classification": "restricted",
                "external_distribution_allowed": "restricted_review_required",
                "last_reviewed": "2026-04-29",
                "access_level": "restricted",
                "language": "en",
            },
            {
                "annex_id": "ANN-003",
                "annex_title": "Classification Marking Appendix",
                "owner": "Records and Security Office",
                "linked_doc_id": "CLASS-ANNEX-B-2025",
                "brief_id": "BRIEF-ANNEX-26",
                "classification": "restricted",
                "external_distribution_allowed": "restricted_review_required",
                "last_reviewed": "2026-04-29",
                "access_level": "restricted",
                "language": "en",
            },
        ),
    ),
}


def generate_synthetic_documents(force: bool = False) -> list[Path]:
    settings = get_settings()
    markdown_dir = settings.generated_corpus_dir / "source_markdown"
    generated_dir = settings.generated_corpus_dir / "generated"
    tables_dir = settings.generated_corpus_dir / "tables"
    for directory in (markdown_dir, generated_dir, tables_dir):
        directory.mkdir(parents=True, exist_ok=True)

    paths: list[Path] = []
    for spec in SPECS:
        markdown_path = markdown_dir / spec.markdown_filename
        if force or not markdown_path.exists():
            markdown_path.write_text(_markdown_for(spec), encoding="utf-8")
        paths.append(markdown_path)

        generated_path = generated_dir / spec.filename
        if force or not generated_path.exists():
            if generated_path.suffix.lower() == ".docx":
                _write_docx(spec, generated_path)
            elif generated_path.suffix.lower() == ".pdf":
                _write_pdf(spec, generated_path)

    table_path = tables_dir / "doctrine_review_tracker.csv"
    if force or not table_path.exists():
        _write_tracker_csv(table_path)
    paths.append(table_path)
    for filename, (headers, rows) in TABLE_DEFINITIONS.items():
        table_path = tables_dir / filename
        if force or not table_path.exists():
            _write_dict_csv(table_path, headers, rows)
        paths.append(table_path)
    return paths


def spec_by_doc_id(doc_id: str) -> SyntheticDocSpec:
    for spec in SPECS:
        if spec.doc_id == doc_id:
            return spec
    raise KeyError(doc_id)


def spec_by_filename(filename: str) -> SyntheticDocSpec:
    stem = Path(filename).stem
    return spec_by_doc_id(stem)


def source_checksum(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _metadata(spec: SyntheticDocSpec, canonical_path: str) -> dict[str, object]:
    return {
        "doc_id": spec.doc_id,
        "title": spec.title,
        "doc_family": spec.doc_family,
        "version": spec.version,
        "status": spec.status,
        "effective_date": spec.effective_date,
        "owner": spec.owner,
        "review_due": spec.review_due,
        "access_level": spec.access_level,
        "language": spec.language,
        "source_type": spec.source_type,
        "canonical_source": canonical_path,
        "allowed_roles": _allowed_roles(spec.access_level),
        "created_at": "2026-05-06T00:00:00Z",
        "authoritative_rank": spec.authoritative_rank,
        "supersedes": list(spec.supersedes),
        "superseded_by": list(spec.superseded_by),
        "cross_references": list(spec.cross_references),
        "applies_to": list(spec.applies_to),
        "not_applicable_to": list(spec.not_applicable_to),
    }


def _markdown_for(spec: SyntheticDocSpec) -> str:
    metadata = yaml.safe_dump(_metadata(spec, f"generated/{spec.filename}"), sort_keys=False).strip()
    sections: list[str] = []
    for section_id, (section, page, body) in enumerate(spec.sections, start=1):
        sections.append(f"## {section}\n\nPage: {page}\nSection ID: S{section_id}\n\n{body.strip()}\n")
    return f"---\n{metadata}\n---\n\n# {spec.title}\n\n" + "\n".join(sections)


def _allowed_roles(access_level: str) -> list[str]:
    if access_level == "restricted":
        return ["planning_lead", "admin"]
    return ["planning_analyst", "planning_lead", "auditor", "admin"]


def _write_docx(spec: SyntheticDocSpec, path: Path) -> None:
    doc = DocxDocument()
    doc.add_heading(spec.title, level=0)
    doc.add_paragraph(f"Document ID: {spec.doc_id}")
    doc.add_paragraph(f"Status: {spec.status}")
    doc.add_paragraph(f"Access level: {spec.access_level}")
    doc.add_paragraph(f"Effective date: {spec.effective_date}")
    for section, page, body in spec.sections:
        doc.add_heading(section, level=1)
        doc.add_paragraph(f"Page: {page}")
        doc.add_paragraph(body)
    doc.save(path)


def _write_pdf(spec: SyntheticDocSpec, path: Path) -> None:
    pdf = canvas.Canvas(str(path), pagesize=letter)
    width, height = letter
    y = height - 72
    pdf.setFont("Helvetica-Bold", 14)
    pdf.drawString(72, y, spec.title[:80])
    y -= 24
    pdf.setFont("Helvetica", 9)
    pdf.drawString(72, y, f"Document ID: {spec.doc_id} | Status: {spec.status} | Access: {spec.access_level}")
    y -= 28
    for section, page, body in spec.sections:
        pdf.setFont("Helvetica-Bold", 11)
        pdf.drawString(72, y, f"SECTION: {section} | PAGE: {page}")
        y -= 16
        pdf.setFont("Helvetica", 9)
        for line in _wrap(body, width=92):
            if y < 72:
                pdf.showPage()
                y = height - 72
            pdf.drawString(72, y, line)
            y -= 13
        y -= 12
    pdf.save()


def _write_tracker_csv(path: Path) -> None:
    headers = (
        "doc_id",
        "title",
        "doc_family",
        "owner",
        "status",
        "effective_date",
        "next_review_due",
        "access_level",
        "language",
    )
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(headers)
        writer.writerows(TRACKER_ROWS)


def _write_dict_csv(path: Path, headers: tuple[str, ...], rows: tuple[dict[str, str], ...]) -> None:
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(headers), lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def _wrap(text: str, width: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    current: list[str] = []
    for word in words:
        proposed = " ".join([*current, word])
        if len(proposed) > width and current:
            lines.append(" ".join(current))
            current = [word]
        else:
            current.append(word)
    if current:
        lines.append(" ".join(current))
    return lines
