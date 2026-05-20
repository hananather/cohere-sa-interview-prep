from __future__ import annotations

import hashlib
import json
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path
from typing import Any

import fitz
import yaml

from defence_agent.config import get_settings


DEFAULT_ALLOWED_ROLES = ["all"]


@dataclass(frozen=True)
class DocumentPage:
    """One normalized source page ready for retrieval and traceability."""

    page_id: str
    doc_id: str
    page_number: int
    text: str
    image_data_url: str
    metadata: dict[str, Any]


def load_document_pages(corpus_dir: Path | None = None) -> list[DocumentPage]:
    """Load the corpus as page-level PDF artifacts.

    The retrieval unit is intentionally a page, not a hand-tuned text chunk.
    Each page can be embedded by Cohere Embed v4 as a rendered page image.
    """

    root = corpus_dir or get_settings().corpus_dir
    manifest_path = root / "manifest.yaml"
    if not manifest_path.exists():
        raise FileNotFoundError(f"Corpus manifest does not exist: {manifest_path}")
    settings = get_settings()
    stat = manifest_path.stat()
    cached_pages = _load_manifest_pages_cached(
        str(manifest_path.resolve()),
        stat.st_mtime_ns,
        stat.st_size,
        settings.parser_backend,
        settings.compass_parser_url,
    )
    return [_copy_page(page) for page in cached_pages]


@lru_cache(maxsize=8)
def _load_manifest_pages_cached(
    manifest_path: str,
    manifest_mtime_ns: int,
    manifest_size: int,
    parser_backend: str,
    compass_parser_url: str,
) -> tuple[DocumentPage, ...]:
    # The mtime and size arguments are part of the cache key.
    _ = (manifest_mtime_ns, manifest_size, parser_backend, compass_parser_url)
    return tuple(_load_manifest_pages(Path(manifest_path)))


def _copy_page(page: DocumentPage) -> DocumentPage:
    return DocumentPage(
        page_id=page.page_id,
        doc_id=page.doc_id,
        page_number=page.page_number,
        text=page.text,
        image_data_url=page.image_data_url,
        metadata=dict(page.metadata),
    )


def _load_manifest_pages(manifest_path: Path) -> list[DocumentPage]:
    manifest = yaml.safe_load(manifest_path.read_text(encoding="utf-8")) or {}
    documents = manifest.get("documents", [])
    if not isinstance(documents, list):
        raise ValueError(f"{manifest_path} must contain a documents list")

    pages: list[DocumentPage] = []
    for entry in documents:
        if not isinstance(entry, dict):
            raise ValueError(f"{manifest_path} contains a non-object document entry")
        metadata = _manifest_metadata(entry)
        doc_id = str(metadata["doc_id"])
        source_path = _resolve_manifest_path(manifest_path.parent, entry)
        pages.extend(_pages_from_source(source_path, metadata, manifest_path=manifest_path))
        if not pages or pages[-1].doc_id != doc_id:
            raise ValueError(f"Manifest entry {doc_id} produced no pages")
    return pages


def _resolve_manifest_path(corpus_root: Path, entry: dict[str, Any]) -> Path:
    raw_path = entry.get("path") or entry.get("pdf_path") or entry.get("docx_path")
    if not raw_path:
        raise ValueError("Manifest document entry is missing path")
    source_path = Path(str(raw_path)).expanduser()
    if not source_path.is_absolute():
        source_path = corpus_root / source_path
    if not source_path.exists():
        raise FileNotFoundError(f"Manifest source file does not exist: {source_path}")
    return source_path


def _manifest_metadata(entry: dict[str, Any]) -> dict[str, Any]:
    metadata = dict(entry.get("metadata") or {})
    for key, value in entry.items():
        if key in {"path", "pdf_path", "docx_path", "metadata"}:
            continue
        metadata.setdefault(key, value)
    if not metadata.get("doc_id"):
        raise ValueError("Manifest document entry is missing doc_id")
    metadata.setdefault("access_level", "unclassified")
    metadata.setdefault("allowed_roles", DEFAULT_ALLOWED_ROLES)
    metadata.setdefault("source_type", "pdf")
    metadata.setdefault("status", "approved")
    metadata.setdefault("language", "en")
    return metadata


def metadata_for_vector_store(metadata: dict[str, Any]) -> dict[str, str | int | float | bool]:
    """Normalize metadata to Chroma-compatible primitive values."""

    normalized: dict[str, str | int | float | bool] = {}
    for key, value in metadata.items():
        if isinstance(value, (str, int, float, bool)):
            normalized[key] = value
        elif value is None:
            normalized[key] = ""
        else:
            normalized[key] = json.dumps(value)
    return normalized


def _pages_from_source(
    source_path: Path,
    frontmatter: dict[str, Any],
    manifest_path: Path | None = None,
) -> list[DocumentPage]:
    settings = get_settings()
    if settings.parser_backend == "compass":
        return _pages_from_compass(source_path, frontmatter, manifest_path=manifest_path)
    if source_path.suffix.lower() == ".pdf":
        return _pages_from_pdf(source_path, frontmatter, manifest_path=manifest_path)
    if source_path.suffix.lower() == ".docx":
        return _pages_from_docx(source_path, frontmatter, manifest_path=manifest_path)
    raise ValueError(f"Unsupported local source format: {source_path.suffix or source_path}")


def _pages_from_pdf(
    pdf_path: Path,
    frontmatter: dict[str, Any],
    manifest_path: Path | None = None,
) -> list[DocumentPage]:
    doc_id = str(frontmatter.get("doc_id") or pdf_path.stem)
    pages: list[DocumentPage] = []
    with fitz.open(pdf_path) as document:
        for page_index, page in enumerate(document, start=1):
            text = page.get_text("text").strip()
            image_bytes = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False).tobytes("png")
            image_data_url = "data:image/png;base64," + _b64(image_bytes)
            metadata = _base_metadata(frontmatter)
            metadata.update(
                {
                    "doc_id": doc_id,
                    "page": page_index,
                    "row_id": "",
                    "parser_backend": "local",
                    "source_pdf_path": str(pdf_path),
                    "manifest_path": str(manifest_path or ""),
                    "page_image_sha256": hashlib.sha256(image_bytes).hexdigest(),
                }
            )
            pages.append(
                DocumentPage(
                    page_id=f"{doc_id}_page_{page_index:03d}",
                    doc_id=doc_id,
                    page_number=page_index,
                    text=text,
                    image_data_url=image_data_url,
                    metadata=metadata,
                )
            )
    return pages


def _pages_from_docx(
    docx_path: Path,
    frontmatter: dict[str, Any],
    manifest_path: Path | None = None,
) -> list[DocumentPage]:
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - exercised only without optional dependency
        raise RuntimeError("python-docx is required to parse DOCX sources locally") from exc

    document = Document(str(docx_path))
    blocks = _docx_blocks(document)
    logical_pages = _logical_docx_pages(blocks)
    doc_id = str(frontmatter.get("doc_id") or docx_path.stem)
    pages: list[DocumentPage] = []
    for page_index, text in enumerate(logical_pages, start=1):
        metadata = _base_metadata(frontmatter)
        metadata.update(
            {
                "doc_id": doc_id,
                "page": page_index,
                "row_id": "",
                "parser_backend": "local",
                "source_docx_path": str(docx_path),
                "manifest_path": str(manifest_path or ""),
                "page_image_sha256": "",
                "logical_page": True,
            }
        )
        pages.append(
            DocumentPage(
                page_id=f"{doc_id}_page_{page_index:03d}",
                doc_id=doc_id,
                page_number=page_index,
                text=text,
                image_data_url="",
                metadata=metadata,
            )
        )
    return pages


def _docx_blocks(document: Any) -> list[str]:
    blocks: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            blocks.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))
    return blocks or [""]


def _logical_docx_pages(blocks: list[str], *, max_chars: int = 3200) -> list[str]:
    pages: list[str] = []
    current: list[str] = []
    current_chars = 0
    for block in blocks:
        block_chars = len(block)
        if current and current_chars + block_chars > max_chars:
            pages.append("\n\n".join(current).strip())
            current = []
            current_chars = 0
        current.append(block)
        current_chars += block_chars + 2
    if current:
        pages.append("\n\n".join(current).strip())
    return pages or [""]


def _pages_from_compass(
    source_path: Path,
    frontmatter: dict[str, Any],
    manifest_path: Path | None = None,
) -> list[DocumentPage]:
    settings = get_settings()
    try:
        from cohere_compass.clients.parser import CompassParserClient
    except ImportError as exc:  # pragma: no cover - exercised only without optional dependency
        raise RuntimeError("cohere-compass-sdk is required when DEFENCE_AGENT_PARSER_BACKEND=compass") from exc

    doc_id = str(frontmatter.get("doc_id") or source_path.stem)
    with CompassParserClient(
        parser_url=settings.compass_parser_url,
        bearer_token=settings.compass_bearer_token or None,
    ) as client:
        documents = client.process_file(filename=str(source_path), file_id=doc_id)
    pages = _pages_from_compass_documents(
        documents,
        source_path=source_path,
        frontmatter=frontmatter,
        manifest_path=manifest_path,
    )
    if not pages:
        raise ValueError(f"Compass parser produced no pages for {source_path}")
    return pages


def _pages_from_compass_documents(
    documents: list[Any],
    *,
    source_path: Path,
    frontmatter: dict[str, Any],
    manifest_path: Path | None,
) -> list[DocumentPage]:
    doc_id = str(frontmatter.get("doc_id") or source_path.stem)
    page_texts: dict[int, list[str]] = {}
    for document in documents:
        chunks = list(getattr(document, "chunks", []) or [])
        if chunks:
            for chunk in chunks:
                page_number = _compass_chunk_page(chunk) or len(page_texts) + 1
                text = _compass_content_text(getattr(chunk, "content", {}) or {})
                if text:
                    page_texts.setdefault(page_number, []).append(text)
            continue
        text = str(getattr(document, "markdown", "") or "")
        if not text:
            text = _compass_content_text(getattr(document, "content", {}) or {})
        if text:
            page_texts.setdefault(1, []).append(text)

    pages: list[DocumentPage] = []
    for page_index in sorted(page_texts):
        metadata = _base_metadata(frontmatter)
        metadata.update(
            {
                "doc_id": doc_id,
                "page": page_index,
                "row_id": "",
                "parser_backend": "compass",
                "source_path": str(source_path),
                "manifest_path": str(manifest_path or ""),
                "page_image_sha256": "",
            }
        )
        pages.append(
            DocumentPage(
                page_id=f"{doc_id}_page_{page_index:03d}",
                doc_id=doc_id,
                page_number=page_index,
                text="\n\n".join(page_texts[page_index]).strip(),
                image_data_url="",
                metadata=metadata,
            )
        )
    return pages


def _compass_chunk_page(chunk: Any) -> int | None:
    origin = getattr(chunk, "origin", None) or {}
    if not isinstance(origin, dict):
        return None
    for key in ("page", "page_number", "page_num"):
        value = origin.get(key)
        try:
            return int(value)
        except (TypeError, ValueError):
            continue
    return None


def _compass_content_text(content: dict[str, Any]) -> str:
    if not isinstance(content, dict):
        return str(content or "")
    preferred = ["text", "markdown", "page_content", "content"]
    values = [str(content.get(key, "") or "") for key in preferred if content.get(key)]
    if not values:
        values = [str(value) for value in content.values() if isinstance(value, str)]
    return "\n".join(value.strip() for value in values if value.strip())


def _base_metadata(frontmatter: dict[str, Any]) -> dict[str, Any]:
    metadata = dict(frontmatter)
    defaults: dict[str, Any] = {
        "title": "",
        "doc_family": "",
        "version": "",
        "status": "approved",
        "effective_date": "",
        "owner": "",
        "review_due": "",
        "access_level": "unclassified",
        "language": "en",
        "source_type": "pdf",
        "source_format": "pdf",
        "normalized_format": "pdf",
        "normalization_method": "",
        "parser_backend": "local",
        "allowed_roles": DEFAULT_ALLOWED_ROLES,
        "authoritative_rank": "",
        "supersedes": [],
        "superseded_by": "",
        "cross_references": [],
        "applies_to": [],
        "not_applicable_to": [],
    }
    for key, value in defaults.items():
        metadata.setdefault(key, value)
    for key in (
        "title",
        "doc_family",
        "version",
        "status",
        "effective_date",
        "owner",
        "review_due",
        "access_level",
        "language",
        "source_type",
        "source_format",
        "normalized_format",
        "normalization_method",
        "parser_backend",
        "authoritative_rank",
        "superseded_by",
    ):
        metadata[key] = str(metadata.get(key, ""))
    return metadata


def _b64(value: bytes) -> str:
    import base64

    return base64.b64encode(value).decode("utf-8")
